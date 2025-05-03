Abslam Omari, [03/05/2025 07:13]
# === 1. INSTALL DEPENDENCIES ===
!pip install pennylane==0.34.0 jax==0.4.28 jaxlib==0.4.28
!pip install transformers torch sentencepiece scikit-learn flask spacy pytesseract pdf2image python-docx
!python -m spacy download fr_core_news_sm
!sudo apt install tesseract-ocr-fra  # For French OCR

# === 2. IMPORTS ===
import os
import numpy as np
import pandas as pd
from pdf2image import convert_from_path
import pytesseract
from docx import Document
import spacy
from transformers import T5ForConditionalGeneration, T5Tokenizer
import torch
import pennylane as qml
from sklearn.cluster import KMeans
from flask import Flask, request, jsonify

# === 3. DATA PROCESSING ===
class ExamProcessor:
    def init(self, dataset_path):
        self.df = pd.read_csv(os.path.join(dataset_path, "metadata.csv"))
        self.nlp = spacy.load("fr_core_news_sm")
        
    def extract_text(self, pdf_path):
        """Convert PDF exam papers to text using OCR"""
        images = convert_from_path(pdf_path)
        text = ""
        for img in images:
            text += pytesseract.image_to_string(img, lang='fra')
        return text
    
    def get_exam_content(self, level=None, subject=None):
        """Filter and process exams by level/subject"""
        filtered = self.df
        if level:
            filtered = filtered[filtered['level'] == level]
        if subject:
            filtered = filtered[filtered['subject'] == subject]
            
        contents = []
        for _, row in filtered.iterrows():
            text = self.extract_text(os.path.join(dataset_path, row['filename']))
            contents.append({
                'level': row['level'],
                'subject': row['subject'],
                'text': text
            })
        return contents

# === 4. QUANTUM NLP SYSTEM ===
class QuantumQuestionGenerator:
    def init(self):
        self.tokenizer = T5Tokenizer.from_pretrained("t5-small")
        self.model = T5ForConditionalGeneration.from_pretrained("t5-small")
        self.nlp = spacy.load("fr_core_news_sm")
        self.dev = qml.device("default.qubit", wires=4)
        
    @qml.qnode(self.dev)
    def quantum_circuit(self, x):
        x = x / np.linalg.norm(x)
        qml.AngleEmbedding(x[:4], wires=range(4))
        qml.BasicEntanglerLayers(weights=np.ones((1,4)), wires=range(4))
        return qml.state()
    
    def quantum_kernel(self, x1, x2):
        state1 = self.quantum_circuit(x1)
        state2 = self.quantum_circuit(x2)
        return np.abs(np.dot(np.conj(state1), state2))**2
    
    def generate_questions(self, text, num_questions=5):
        input_text = f"générer des questions en français: {text}"
        input_ids = self.tokenizer.encode(input_text, return_tensors="pt")
        
        outputs = self.model.generate(
            input_ids,
            max_length=100,
            num_return_sequences=num_questions,
            do_sample=True,
            temperature=0.7
        )
        return [self.tokenizer.decode(output, skip_special_tokens=True) 
               for output in outputs]
    
    def optimize_diversity(self, questions, n_clusters=3):
        features = []
        for q in questions:
            doc = self.nlp(q)
            features.append([
                len(q),
                len(doc),
                int(any(t.pos_ == "VERB" for t in doc)),
                int(any(t.text.lower() in ['qui','quoi','où','comment','pourquoi'] for t in doc))
            ])
        
        # Quantum-enhanced clustering
        kernel = np.zeros((len(features), len(features)))
        for i in range(len(features)):
            for j in range(i, len(features)):
                kernel[i,j] = self.quantum_kernel(np.array(features[i]), np.array(features[j]))
                kernel[j,i] = kernel[i,j]
                
        kmeans = KMeans(n_clusters=min(n_clusters, len(questions)))
        clusters = kmeans.fit_predict(kernel)
        
        return [questions[np.where(clusters == i)[0][0]] for i in range(kmeans.n_clusters)]

Abslam Omari, [03/05/2025 07:13]
# === 5. APPLICATION PIPELINE ===
def process_algerian_exams(dataset_path, output_dir):
    processor = ExamProcessor(dataset_path)
    generator = QuantumQuestionGenerator()
    
    # Focus on BAC level exams
    bac_exams = processor.get_exam_content(level="شهادة البكالوريا")
    
    for exam in bac_exams:
        print(f"\nProcessing {exam['subject']} exam...")
        questions = generator.generate_questions(exam['text'])
        optimized = generator.optimize_diversity(questions)
        
        # Save to Word document
        doc = Document()
        doc.add_heading(f"Questions générées - {exam['subject']}", level=1)
        for i, q in enumerate(optimized):
            doc.add_paragraph(f"{i+1}. {q}")
        doc.save(os.path.join(output_dir, f"{exam['subject']}_questions.docx"))

# === 6. FLASK API ===
app = Flask(name)
generator = QuantumQuestionGenerator()

@app.route('/generate', methods=['POST'])
def generate_questions():
    data = request.json
    arabic_text = data.get('text', '')
    
    # Simple Arabic to French translation placeholder
    # In production, use proper translation API
    french_text = arabic_text.replace("الرياضيات", "Mathématiques").replace("العربية", "Arabe")
    
    questions = generator.generate_questions(french_text)
    optimized = generator.optimize_diversity(questions)
    
    return jsonify({
        'original_text': arabic_text,
        'questions': optimized
    })

# === 7. EXECUTION ===
if name == 'main':
    # Example usage with Kaggle dataset
    dataset_path = "/kaggle/input/algerian-education-exam-papers-dataset"
    output_dir = "/kaggle/working/generated_questions"
    os.makedirs(output_dir, exist_ok=True)
    
    # Process all BAC exams
    process_algerian_exams(dataset_path, output_dir)
    
    # Start API (uncomment for deployment)
    # app.run(host='0.0.0.0', port=5000)
